"""Input loading and workbook validation."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Iterable

import numpy as np
import pandas as pd
from openpyxl import load_workbook


REQUIRED_SHEETS = ("Main", "Prices", "Turnover")

# The source workbook has duplicate column labels, so positions are used to give
# each of the first 20 columns a stable and unambiguous name.
MAIN_COLUMNS = [
    "Ticker",
    "Name",
    "PriceUSD",
    "ADV20USD_raw",
    "MarketCapUSD_raw",
    "MarketCapLocal_raw",
    "PriceLocal",
    "PE",
    "UnusedColumnI",
    "Sector",
    "NextEarning",
    "Vol20D",
    "AvgVol5D",
    "AvgVol20D",
    "Country",
    "Return5D",
    "Return3M",
    "EPSNow",
    "EPS30d",
    "EPSChangeRaw",
]


def parse_scaled_number(value: object) -> float:
    """Convert Bloomberg-style strings such as ``43.75M`` or ``21.52B``.

    Parameters
    ----------
    value:
        A numeric value or a string ending in K, M, B, or T.

    Returns
    -------
    float
        Parsed numeric value, or ``numpy.nan`` when parsing fails.
    """

    if value is None:
        return np.nan
    if isinstance(value, (int, float, np.integer, np.floating)):
        return float(value) if not pd.isna(value) else np.nan

    text = str(value).strip().replace(",", "")
    match = re.fullmatch(
        r"([+-]?(?:\d+(?:\.\d*)?|\.\d+))\s*([KMBT]?)",
        text,
        flags=re.IGNORECASE,
    )
    if match is None:
        return np.nan

    scale = {
        "": 1.0,
        "K": 1_000.0,
        "M": 1_000_000.0,
        "B": 1_000_000_000.0,
        "T": 1_000_000_000_000.0,
    }[match.group(2).upper()]
    return float(match.group(1)) * scale


def validate_workbook(path: Path) -> None:
    """Validate that the input workbook contains the required sheets."""

    if not path.exists():
        raise FileNotFoundError(f"Input workbook not found: {path}")
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("The input file must be an .xlsx or .xlsm workbook.")

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        missing = [name for name in REQUIRED_SHEETS if name not in workbook.sheetnames]
    finally:
        workbook.close()

    if missing:
        raise ValueError(
            "Input workbook is missing required sheets: " + ", ".join(missing)
        )


def infer_as_of_date(path: Path, override: str | pd.Timestamp | None = None) -> pd.Timestamp:
    """Resolve the model's as-of date.

    Priority:
    1. A user-supplied date.
    2. The workbook's ``modified`` metadata date.
    3. The file-system modification date.

    The event score depends on this date, so production runs should normally
    pass ``--as-of-date YYYY-MM-DD`` explicitly.
    """

    if override is not None:
        parsed = pd.Timestamp(override)
        if pd.isna(parsed):
            raise ValueError(f"Invalid as-of date: {override}")
        return parsed.normalize()

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        modified = workbook.properties.modified
    finally:
        workbook.close()

    if modified is not None:
        return pd.Timestamp(modified.date())

    return pd.Timestamp(path.stat().st_mtime, unit="s").normalize()


def _clean_tickers(series: pd.Series) -> pd.Series:
    return series.astype("string").str.strip()


def load_main_sheet(path: Path) -> pd.DataFrame:
    """Load and clean the first 20 columns of the ``Main`` sheet."""

    data = pd.read_excel(
        path,
        sheet_name="Main",
        usecols="A:T",
        engine="openpyxl",
        dtype=object,
    )
    if data.shape[1] != len(MAIN_COLUMNS):
        raise ValueError(
            f"Main sheet should expose {len(MAIN_COLUMNS)} columns A:T; "
            f"found {data.shape[1]}."
        )

    data.columns = MAIN_COLUMNS
    data["Ticker"] = _clean_tickers(data["Ticker"])

    # This removes the summary line and the blank separator rows without relying
    # on a fixed row number.
    data = data[data["Ticker"].notna() & data["Ticker"].ne("")].copy()
    data = data.drop_duplicates(subset="Ticker", keep="first").reset_index(drop=True)

    data["ADV20USD"] = data["ADV20USD_raw"].map(parse_scaled_number)
    data["MarketCapUSD"] = data["MarketCapUSD_raw"].map(parse_scaled_number)
    data["MarketCapLocal"] = data["MarketCapLocal_raw"].map(parse_scaled_number)

    numeric_columns = [
        "PriceUSD",
        "PriceLocal",
        "PE",
        "Vol20D",
        "AvgVol5D",
        "AvgVol20D",
        "Return5D",
        "Return3M",
        "EPSNow",
        "EPS30d",
        "EPSChangeRaw",
    ]
    for column in numeric_columns:
        data[column] = pd.to_numeric(data[column], errors="coerce")

    data["Sector"] = data["Sector"].astype("string").str.strip()
    data["Sector"] = data["Sector"].replace("", pd.NA).fillna("Unknown")
    data["Country"] = data["Country"].astype("string").str.strip()
    data["Name"] = data["Name"].astype("string").str.strip()

    return data


def _ordered_time_columns(columns: Iterable[object]) -> list[str]:
    """Return time columns ordered t-0, t-1, t-2, ... ."""

    parsed: list[tuple[int, str]] = []
    for column in columns:
        text = str(column).strip()
        match = re.fullmatch(r"t-(\d+)", text)
        if match:
            parsed.append((int(match.group(1)), text))
    return [name for _, name in sorted(parsed)]


def load_panel_sheet(path: Path, sheet_name: str) -> pd.DataFrame:
    """Load a ticker-by-time panel from ``Prices`` or ``Turnover``."""

    panel = pd.read_excel(path, sheet_name=sheet_name, engine="openpyxl")
    if panel.empty or panel.shape[1] < 2:
        raise ValueError(f"Sheet {sheet_name!r} does not contain a usable panel.")

    panel = panel.rename(columns={panel.columns[0]: "Ticker"})
    panel["Ticker"] = _clean_tickers(panel["Ticker"])
    panel = panel[panel["Ticker"].notna() & panel["Ticker"].ne("")].copy()
    panel = panel.drop_duplicates(subset="Ticker", keep="first")

    time_columns = _ordered_time_columns(panel.columns[1:])
    if "t-0" not in time_columns:
        raise ValueError(f"Sheet {sheet_name!r} must contain a t-0 column.")

    panel = panel[["Ticker", *time_columns]]
    for column in time_columns:
        panel[column] = pd.to_numeric(panel[column], errors="coerce")

    return panel.set_index("Ticker")


def load_source_data(path: Path) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """Load all three required input sheets."""

    validate_workbook(path)
    main = load_main_sheet(path)
    prices = load_panel_sheet(path, "Prices")
    turnover = load_panel_sheet(path, "Turnover")
    return main, prices, turnover
