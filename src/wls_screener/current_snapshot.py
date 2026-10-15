"""Replay the October 2026 five-workbook WLS snapshot without private data.

This module deliberately uses the last *completed* close (``Price:D-1``) from
the current export. ``Price`` is an unsynchronised intraday field and is never
used to rank a stock. The October 14 export supplies USD price, cap and ADV;
older local-currency price panels are joined through an inferred FX ratio.
"""

from __future__ import annotations

import argparse
import csv
import json
import math
from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

import numpy as np
import pandas as pd
from dateutil.relativedelta import relativedelta
from openpyxl import load_workbook


# Bloomberg changed these ticker codes between the historical and current
# exports. Mapping is deliberately small and explicit; ambiguous joins fail.
HISTORICAL_TO_CURRENT = {
    "TCBI UW Equity": "TCBI VT Equity",
    "PFSE GY Equity": "PFS0 GY Equity",
    "EAND DH Equity": "ETISALAT DH Equity",
    "OBK UN Equity": "OBK VT Equity",
}


def number(value: object) -> float:
    try:
        result = float(value)
    except (TypeError, ValueError):
        return math.nan
    return result if math.isfinite(result) else math.nan


def valid_positive(value: object) -> bool:
    value = number(value)
    return math.isfinite(value) and value > 0


def signal_close_usd(current_row: dict) -> float:
    """Use the completed D-1 close, never the mixed-time ``Price`` field."""
    return number(current_row["Price:D-1"])


def investable_gates(
    *, close_usd: float, adv_usd: float, cap_usd: float,
    sector: str, has_history: bool,
) -> dict[str, bool]:
    """Apply the strict WLS-to-investable-universe screen before scoring."""
    return {
        "price_gate": math.isfinite(close_usd) and close_usd > 1.0,
        "adv_gate": math.isfinite(adv_usd) and adv_usd > 15_000_000.0,
        "cap_gate": math.isfinite(cap_usd) and cap_usd > 500_000_000.0,
        "sector_gate": str(sector).strip().casefold()
        not in {"", "--", "unknown", "nan", "none", "n/a", "na"},
        "history_gate": has_history,
    }


def date_value(value: object) -> date | None:
    """Parse Excel dates and Bloomberg YYYYMMDD expected-report dates."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if value is None or str(value).strip() in {"", "--"}:
        return None
    text = str(value).strip()
    try:
        return datetime.strptime(str(int(float(text))), "%Y%m%d").date()
    except (ValueError, TypeError, OverflowError):
        return None


def read_worksheet(path: Path, columns: tuple[str, ...]) -> dict[str, dict]:
    """Read a Bloomberg Worksheet, tolerating its trailing newlines in headers."""
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        if "Worksheet" not in book:
            raise ValueError(f"{path} has no Worksheet sheet")
        rows = book["Worksheet"].iter_rows(values_only=True)
        header = [str(item).strip() if item is not None else "" for item in next(rows)]
        missing = sorted(set(columns) - set(header))
        if missing:
            raise ValueError(f"{path} is missing columns: {missing}")
        positions = {name: header.index(name) for name in columns}
        result: dict[str, dict] = {}
        for row in rows:
            ticker = str(row[positions["Ticker"]] or "").strip()
            if ticker:
                if ticker in result:
                    raise ValueError(f"Duplicate ticker {ticker!r} in {path}")
                result[ticker] = {name: row[i] for name, i in positions.items()}
        return result
    finally:
        book.close()


def read_dated_panel(path: Path, sheet: str) -> tuple[list[date], dict[str, np.ndarray]]:
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet not in book:
            raise ValueError(f"{path} has no {sheet!r} sheet")
        rows = book[sheet].iter_rows(values_only=True)
        header = next(rows)
        dated = sorted(
            (item.date(), i) for i, item in enumerate(header) if isinstance(item, datetime)
        )
        dates = [day for day, _ in dated]
        if not dates or len(set(dates)) != len(dates):
            raise ValueError(f"{sheet} has missing or duplicate date headers")
        result = {}
        for row in rows:
            ticker = str(row[0] or "").strip()
            if ticker:
                if ticker in result:
                    raise ValueError(f"Duplicate ticker {ticker!r} in {sheet}")
                result[ticker] = np.asarray([number(row[i]) for _, i in dated], dtype=float)
        return dates, result
    finally:
        book.close()


def current_key_map(old_tickers: object) -> dict[str, str]:
    """Map an older ticker set to current symbols and reject alias collisions."""
    mapped = {}
    for old in old_tickers:
        current = HISTORICAL_TO_CURRENT.get(old, old)
        if current in mapped:
            raise ValueError(f"Ticker alias collision: {mapped[current]!r} and {old!r}")
        mapped[current] = old
    return mapped


def infer_closed_exchanges(
    pairs: list[tuple[str, object, object]], *,
    minimum_pairs: int = 5, same_threshold: float = 0.95,
    abs_tolerance: float = 1e-7,
) -> tuple[set[str], list[dict]]:
    """Infer an exchange closure from widespread quote carry-forward.

    One unchanged stock is ordinary; >=95% unchanged among at least five
    quoted names on the same exchange is evidence of a closed session.
    """
    counts: dict[str, dict[str, int]] = defaultdict(lambda: {"paired": 0, "unchanged": 0})
    for exchange, earlier, later in pairs:
        if valid_positive(earlier) and valid_positive(later):
            counts[exchange]["paired"] += 1
            counts[exchange]["unchanged"] += int(
                math.isclose(number(earlier), number(later), rel_tol=1e-6,
                             abs_tol=abs_tolerance)
            )
    closed = {exchange for exchange, item in counts.items()
              if item["paired"] >= minimum_pairs
              and item["unchanged"] / item["paired"] >= same_threshold}
    audit = [{"exchange": exchange, **item,
              "unchanged_share": item["unchanged"] / item["paired"] if item["paired"] else None,
              "closed_inferred": exchange in closed}
             for exchange, item in sorted(counts.items())]
    return closed, audit


def sector_percentile(series: pd.Series) -> pd.Series:
    good = series.dropna()
    result = pd.Series(0.5, index=series.index, dtype=float)
    if len(good):
        result.loc[good.index] = (good.rank(method="average") - 0.5) / len(good)
    return result


def _write_csv(path: Path, rows: list[dict]) -> None:
    if not rows:
        raise ValueError(f"Refusing to write an empty CSV: {path}")
    columns = list(dict.fromkeys(key for row in rows for key in row))
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def _records(frame: pd.DataFrame) -> list[dict]:
    rows = []
    for row in frame.to_dict("records"):
        clean = {}
        for key, value in row.items():
            if isinstance(value, (float, np.floating)):
                clean[key] = round(float(value), 8) if math.isfinite(value) else None
            elif isinstance(value, np.integer):
                clean[key] = int(value)
            else:
                clean[key] = value
        rows.append(clean)
    return rows


def _recent_business_days(signal_date: date) -> tuple[date, date, date]:
    """Dates represented by the bridge D-2/D-1 and current D-2 columns.

    This replay expects the five-workbook sequence: an old panel through
    Wednesday, a Friday sector export, Monday interim data and Tuesday close.
    Individual market holidays are handled by closure inference below.
    """
    interim = date.fromisoformat(str(np.busday_offset(str(signal_date), -1)))
    bridge_later = date.fromisoformat(str(np.busday_offset(str(interim), -1)))
    bridge_earlier = date.fromisoformat(str(np.busday_offset(str(interim), -2)))
    return bridge_earlier, bridge_later, interim


def _catalyst(event: date | None, decision_date: date) -> tuple[float, int | None]:
    if event is None or event < decision_date:
        return 0.0, None
    days = int(np.busday_count(np.datetime64(decision_date), np.datetime64(event)))
    score = 1.0 if days <= 5 else 0.75 if days <= 10 else 0.5 if days <= 15 else 0.0
    return score, days


def replay(
    *, historical_prices: Path, eps_volume: Path, sector_snapshot: Path,
    interim_snapshot: Path, current_snapshot: Path, output_dir: Path,
    signal_date: date = date(2026, 10, 13), decision_date: date = date(2026, 10, 14),
) -> dict:
    """Reconstruct the 10k-to-5k filter, sector scores and 3A/2B targets."""
    if decision_date <= signal_date:
        raise ValueError("decision_date must be later than the completed signal_date")
    bridge_earlier, bridge_later, interim_date = _recent_business_days(signal_date)
    prices_dates, prices = read_dated_panel(historical_prices, "Prices")
    eps_dates, eps = read_dated_panel(eps_volume, "EPS")
    volume_dates, volumes = read_dated_panel(eps_volume, "Trad Volume")
    if prices_dates != eps_dates or prices_dates != volume_dates:
        raise ValueError("Historical price, EPS and volume date headers differ")
    if set(prices) != set(eps) or set(prices) != set(volumes):
        raise ValueError("Historical price, EPS and volume ticker sets differ")
    if max(prices_dates) >= bridge_earlier:
        raise ValueError("Historical panel overlaps the bridge snapshot")
    if max(prices_dates) < bridge_earlier - timedelta(days=7):
        raise ValueError("Historical panel ends too early for this snapshot sequence")
    old = read_worksheet(sector_snapshot,
                         ("Ticker", "Name", "Price:D-2", "Price:D-1",
                          "GICS Ind Name", "GICS Sector"))
    interim = read_worksheet(interim_snapshot,
                             ("Ticker", "Name", "Price:D-1", "Volume:D-1"))
    current = read_worksheet(current_snapshot,
                             ("Ticker", "Name", "Weight", "Price", "Curncy", "BEst EPS BF12M",
                              "Exp Rep Dt", "Avg D Val Traded 20D:D-20", "Market Cap:D-1",
                              "Volume:D-2", "Volume:D-1", "Price:D-2", "Price:D-1"))
    hist_alias = current_key_map(prices)
    old_alias = current_key_map(old)
    interim_alias = current_key_map(interim)

    # A median USD/local rate for each currency robustly splices the old local
    # price history to the completed USD close, with no double conversion.
    fx_pairs: dict[str, list[float]] = defaultdict(list)
    for ticker, row in current.items():
        previous_ticker = interim_alias.get(ticker)
        if previous_ticker:
            local = number(interim[previous_ticker]["Price:D-1"])
            usd = number(row["Price:D-2"])
            if valid_positive(local) and valid_positive(usd):
                fx_pairs[str(row["Curncy"] or "").strip()].append(usd / local)
    fx = {currency: float(np.median(rates)) for currency, rates in fx_pairs.items()
          if len(rates) >= 2}

    # Compare adjacent quotes across each exchange. The first two comparisons
    # use local values; the third compares local Friday versus local Monday.
    last_hist_index = prices_dates.index(max(prices_dates))
    day8_pairs, day9_pairs, day12_pairs = [], [], []
    for ticker, old_ticker in hist_alias.items():
        older = old_alias.get(ticker)
        prior = interim_alias.get(ticker)
        if not older:
            continue
        exchange = ticker.split()[1] if len(ticker.split()) > 1 else ""
        p7 = prices[old_ticker][last_hist_index]
        p8 = old[older]["Price:D-2"]
        p9 = old[older]["Price:D-1"]
        day8_pairs.append((exchange, p7, p8))
        day9_pairs.append((exchange, p8, p9))
        if prior:
            day12_pairs.append((exchange, p9, interim[prior]["Price:D-1"]))
    closed8, audit8 = infer_closed_exchanges(day8_pairs)
    closed9, audit9 = infer_closed_exchanges(day9_pairs)
    closed12, audit12 = infer_closed_exchanges(day12_pairs, abs_tolerance=1e-4)

    eps_date = max(day for day in eps_dates if day < bridge_earlier)
    baseline_date = max(day for day in eps_dates if day <= eps_date - timedelta(days=30))
    eps_index = eps_dates.index(eps_date)
    baseline_index = eps_dates.index(baseline_date)
    three_month_start = decision_date - relativedelta(months=3)
    filter_rows = []
    scored_rows = []
    for ticker, row in current.items():
        older, prior, historical = (old_alias.get(ticker), interim_alias.get(ticker),
                                    hist_alias.get(ticker))
        sector = str(old[older]["GICS Sector"] or "").strip() if older else ""
        currency = str(row["Curncy"] or "").strip()
        rate = fx.get(currency, math.nan)
        close_usd = signal_close_usd(row)
        adv_usd = number(row["Avg D Val Traded 20D:D-20"])
        cap_usd = number(row["Market Cap:D-1"])
        gates = investable_gates(
            close_usd=close_usd, adv_usd=adv_usd, cap_usd=cap_usd,
            sector=sector,
            has_history=bool(historical and older and prior and valid_positive(rate)),
        )
        eligible = all(gates.values())
        filter_rows.append({"ticker": ticker, "name": row["Name"], "sector": sector,
                            "currency": currency, "signal_close_usd": close_usd,
                            "adv20_usd_d20": adv_usd, "market_cap_usd_d1": cap_usd,
                            **gates, "eligible": eligible})
        if not eligible:
            continue

        historical_px = prices[historical]
        historical_vol = volumes[historical]
        active = np.flatnonzero(np.isfinite(historical_px) & (historical_px > 0) &
                                np.isfinite(historical_vol) & (historical_vol > 0))
        px = [float(historical_px[i]) for i in active]
        days = [prices_dates[i] for i in active]
        sparse_vol = [float(historical_vol[i]) for i in active]
        exchange = ticker.split()[1]
        for day, field, closed in ((bridge_earlier, "Price:D-2", closed8),
                                   (bridge_later, "Price:D-1", closed9)):
            bridge_px = number(old[older][field])
            if exchange not in closed and valid_positive(bridge_px):
                px.append(bridge_px)
                days.append(day)
        p12 = number(interim[prior]["Price:D-1"])
        v12 = number(row["Volume:D-2"])
        if exchange not in closed12 and valid_positive(p12):
            px.append(p12)
            days.append(interim_date)
            if valid_positive(v12):
                sparse_vol.append(v12)
        v13 = number(row["Volume:D-1"])
        if valid_positive(close_usd) and valid_positive(v13):
            px.append(close_usd / rate)
            days.append(signal_date)
            sparse_vol.append(v13)
        if not px:
            continue
        px_array = np.asarray(px)
        last = px_array[-1]
        lookback = [i for i, day in enumerate(days) if day <= three_month_start]
        old_price = px_array[lookback[-1]] if lookback else math.nan
        return3m = last / old_price - 1 if valid_positive(old_price) else math.nan
        return5d = last / px_array[-6] - 1 if len(px_array) >= 6 else math.nan
        ma20 = float(px_array[-20:].mean()) if len(px_array) >= 20 else math.nan
        ma50 = float(px_array[-50:].mean()) if len(px_array) >= 50 else math.nan
        vol20 = (float(np.std(np.diff(np.log(px_array[-20:])), ddof=1) * math.sqrt(260) * 100)
                 if len(px_array) >= 20 else math.nan)
        old_observed_vol = historical_vol[active]
        volume_ratio = (float(old_observed_vol[-5:].mean() / old_observed_vol[-20:].mean())
                        if len(old_observed_vol) >= 20 else math.nan)
        sparse_ratio = (float(np.mean(sparse_vol[-5:]) / np.mean(sparse_vol[-20:]))
                        if len(sparse_vol) >= 20 else math.nan)
        eps_now = eps[historical][eps_index]
        eps_then = eps[historical][baseline_index]
        eps_revision = ((eps_now - eps_then) / (abs(eps_now) + abs(eps_then) + 1e-12)
                        if math.isfinite(eps_now) and math.isfinite(eps_then) else math.nan)
        event = date_value(row["Exp Rep Dt"])
        catalyst, business_days = _catalyst(event, decision_date)
        scored_rows.append({"ticker": ticker, "name": row["Name"], "sector": sector,
                            "industry": old[older]["GICS Ind Name"], "currency": currency,
                            "last_trade_date_inferred": days[-1].isoformat(),
                            "signal_close_usd": close_usd,
                            "adv20_usd_mio_d20": adv_usd / 1e6,
                            "market_cap_usd_mio_d1": cap_usd / 1e6,
                            "return3m": return3m, "return5d": return5d,
                            "eps_revision": eps_revision,
                            "distance20": last / ma20 - 1 if valid_positive(ma20) else math.nan,
                            "distance50": last / ma50 - 1 if valid_positive(ma50) else math.nan,
                            "volume_ratio_primary": volume_ratio,
                            "volume_ratio_sparse": sparse_ratio, "vol20_pct": vol20,
                            "next_earnings": (
                                event.isoformat() if event and event >= decision_date else None
                            ),
                            "business_days_to_earnings": business_days,
                            "catalyst": catalyst,
                            "current_eps_bf12m_unscored": number(row["BEst EPS BF12M"]),
                            "fx_usd_per_local_splice": rate})

    if len(scored_rows) < 5:
        raise ValueError("Fewer than five names have complete eligibility and history")
    frame = pd.DataFrame.from_records(scored_rows)
    for source, target in (("return3m", "rank_3m"), ("return5d", "rank_5d"),
                           ("eps_revision", "rank_eps"), ("distance20", "rank_ma20"),
                           ("distance50", "rank_ma50"),
                           ("volume_ratio_primary", "rank_volume"),
                           ("volume_ratio_sparse", "rank_volume_sparse"),
                           ("vol20_pct", "rank_vol20")):
        frame[target] = (
            frame.groupby("sector", group_keys=False)[source].transform(sector_percentile)
        )
    frame["price_confirmation"] = (.40 * frame.rank_ma20 + .30 * frame.rank_ma50 +
                                    .30 * frame.rank_volume)
    frame["a_score"] = (.40 * frame.rank_3m + .25 * frame.rank_eps +
                        .25 * frame.price_confirmation + .10 * frame.rank_5d)
    frame["b_score"] = (.35 * frame.catalyst + .30 * frame.rank_vol20 +
                        .20 * frame.rank_5d + .15 * frame.rank_volume)
    sparse_confirmation = (.40 * frame.rank_ma20 + .30 * frame.rank_ma50 +
                           .30 * frame.rank_volume_sparse)
    frame["a_score_sparse"] = (.40 * frame.rank_3m + .25 * frame.rank_eps +
                               .25 * sparse_confirmation + .10 * frame.rank_5d)
    frame["b_score_sparse"] = (.35 * frame.catalyst + .30 * frame.rank_vol20 +
                               .20 * frame.rank_5d + .15 * frame.rank_volume_sparse)

    a_keys = ["a_score", "rank_3m", "rank_eps", "adv20_usd_mio_d20", "ticker"]
    a_dirs = [False, False, False, False, True]
    b_keys = ["b_score", "catalyst", "rank_vol20", "rank_5d", "adv20_usd_mio_d20", "ticker"]
    b_dirs = [False, False, False, False, False, True]
    a = frame.sort_values(a_keys, ascending=a_dirs, kind="mergesort")
    b = frame.sort_values(b_keys, ascending=b_dirs, kind="mergesort")
    frame["a_rank"] = frame.ticker.map({ticker: i for i, ticker in enumerate(a.ticker, 1)})
    frame["b_rank"] = frame.ticker.map({ticker: i for i, ticker in enumerate(b.ticker, 1)})
    a3 = a.head(3)
    b2 = b.loc[~b.ticker.isin(a3.ticker)].head(2)
    targets = pd.concat([a3.assign(slot=["A1", "A2", "A3"]),
                         b2.assign(slot=["B1", "B2"])])
    targets["target_weight_pct"] = 20
    a_sparse = frame.sort_values(["a_score_sparse", *a_keys[1:]],
                                 ascending=a_dirs, kind="mergesort")
    b_sparse = frame.sort_values(["b_score_sparse", *b_keys[1:]],
                                 ascending=b_dirs, kind="mergesort")
    a_sparse3 = list(a_sparse.head(3).ticker)
    b_sparse2 = list(b_sparse.loc[~b_sparse.ticker.isin(a_sparse3)].head(2).ticker)

    output_dir.mkdir(parents=True, exist_ok=True)
    _write_csv(output_dir / "filter_audit.csv", filter_rows)
    _write_csv(output_dir / "eligible_scored.csv", _records(frame))
    _write_csv(output_dir / "top20_A.csv", _records(a.head(20).assign(rank=range(1, 21))))
    _write_csv(output_dir / "top20_B.csv", _records(b.head(20).assign(rank=range(1, 21))))
    _write_csv(output_dir / "targets_A3_B2.csv", _records(targets))
    calendar_rows = []
    for day, records in ((bridge_earlier, audit8), (bridge_later, audit9),
                         (interim_date, audit12)):
        calendar_rows.extend({"date": day.isoformat(), **row} for row in records)
    _write_csv(output_dir / "exchange_calendar_inference.csv", calendar_rows)
    summary = {
        "signal_date": signal_date.isoformat(), "decision_date": decision_date.isoformat(),
        "universe_count": len(current),
        "price_pass": sum(r["price_gate"] for r in filter_rows),
        "price_adv_pass": sum(r["price_gate"] and r["adv_gate"] for r in filter_rows),
        "price_adv_cap_pass": sum(r["price_gate"] and r["adv_gate"] and r["cap_gate"]
                                  for r in filter_rows),
        "fully_scored": len(frame),
        "historical_eps_dates": [baseline_date.isoformat(), eps_date.isoformat()],
        "primary_volume_through": eps_date.isoformat(),
        "bridge_dates": [bridge_earlier.isoformat(), bridge_later.isoformat(),
                         interim_date.isoformat()],
        "closed_exchange_suffixes_inferred": {
            bridge_earlier.isoformat(): sorted(closed8),
            bridge_later.isoformat(): sorted(closed9),
            interim_date.isoformat(): sorted(closed12),
        },
        "targets": [{"slot": row["slot"], "ticker": row["ticker"],
                     "target_weight_pct": 20} for row in _records(targets)],
        "sparse_updated_volume_sensitivity": {"A3": a_sparse3, "B2": b_sparse2},
        "limitations": [
            "Current Price is an unsynchronised snapshot; only Price:D-1 enters the signal.",
            "Current BF12M EPS is not combined with older FY1 EPS estimates; "
            "FY1 revisions end at the historical panel cutoff.",
            "Primary 5/20 share-volume ratio ends at the historical panel cutoff; "
            "missing bridge-day volumes affect updated-volume sensitivity.",
            "The supplied 20-day traded-value field is lagged to D-20 rather than "
            "the signal close.",
            "Exchange closures are inferred from quote carry-forward; expected "
            "report dates are supplied estimates, not issuer confirmations.",
            "The five 20% targets are model weights, not executable orders or a backtest P&L.",
        ],
    }
    (output_dir / "summary.json").write_text(
        json.dumps(summary, indent=2, ensure_ascii=False, allow_nan=False) + "\n", encoding="utf-8")
    return summary


def main(argv: list[str] | None = None) -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--historical-prices", required=True, type=Path,
                        help="100-day WLS workbook with Prices sheet, through 2026-10-07")
    parser.add_argument("--eps-volume", required=True, type=Path,
                        help="EPS & Turnover workbook with EPS and Trad Volume sheets")
    parser.add_argument("--sector-snapshot", required=True, type=Path,
                        help="2026-10-10 Worksheet with GICS sectors and October 8/9 prices")
    parser.add_argument("--interim-snapshot", required=True, type=Path,
                        help="2026-10-13 Worksheet with completed October 12 local price")
    parser.add_argument("--current-snapshot", required=True, type=Path,
                        help="2026-10-14 USD Worksheet with Price:D-1 and filter fields")
    parser.add_argument("--output-dir", required=True, type=Path,
                        help="Destination outside the Git repository for private results")
    parser.add_argument("--signal-date", type=date.fromisoformat, default=date(2026, 10, 13))
    parser.add_argument("--decision-date", type=date.fromisoformat, default=date(2026, 10, 14))
    args = parser.parse_args(argv)
    summary = replay(**vars(args))
    headline_keys = ("signal_date", "decision_date", "universe_count", "fully_scored", "targets")
    print(json.dumps({key: summary[key] for key in headline_keys},
                     indent=2, ensure_ascii=False))
