"""CSV, JSON, and formatted Excel exports."""

from __future__ import annotations

import json
from dataclasses import asdict
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .pipeline import ScreeningResult


DARK_BLUE = "17365D"
MEDIUM_BLUE = "2F75B5"
LIGHT_BLUE = "D9EAF7"
LIGHT_GRAY = "F2F2F2"
WHITE = "FFFFFF"
ORANGE = "C65911"
GREEN = "008000"


TOP_A_COLUMNS = [
    "ScreenARank",
    "Ticker",
    "Name",
    "Sector",
    "Country",
    "ScreenAScore",
    "Return3M",
    "RankReturn3M",
    "EPSRevision",
    "RankEPSRevision",
    "DistanceMA20",
    "DistanceMA50",
    "VolumeRatio5D20D",
    "PriceConfirmationA",
    "Return5D",
    "RankReturn5D",
    "NextEarningDate",
    "ADV20USD",
    "PassesOptionalFilter",
]

TOP_B_COLUMNS = [
    "ScreenBRank",
    "Ticker",
    "Name",
    "Sector",
    "Country",
    "ScreenBScore",
    "NextEarningDate",
    "DaysToEarnings",
    "CatalystScore",
    "Vol20D",
    "RankVol20D",
    "Return5D",
    "RankReturn5D",
    "VolumeRatio5D20D",
    "RankVolumeRatio",
    "Return3M",
    "ScreenARank",
    "ADV20USD",
    "PassesOptionalFilter",
]

ALL_COLUMNS = [
    "ScreenARank",
    "ScreenBRank",
    "Ticker",
    "Name",
    "Sector",
    "Country",
    "PriceUSD",
    "MarketCapUSD",
    "ADV20USD",
    "Return3M",
    "RankReturn3M",
    "EPSNow",
    "EPS30d",
    "EPSRevision",
    "RankEPSRevision",
    "Price_t0",
    "MA20",
    "DistanceMA20",
    "RankDistanceMA20",
    "MA50",
    "DistanceMA50",
    "RankDistanceMA50",
    "AvgVol5D",
    "AvgVol20D",
    "VolumeRatio5D20D",
    "RankVolumeRatio",
    "PriceConfirmationA",
    "Return5D",
    "RankReturn5D",
    "ScreenAScore",
    "NextEarningDate",
    "DaysToEarnings",
    "CatalystScore",
    "Vol20D",
    "RankVol20D",
    "ScreenBScore",
    "TurnoverRatio5D20D",
    "PriceHistoryCount",
    "PassesOptionalFilter",
]

DISPLAY_NAMES = {
    "ScreenARank": "Screen A Rank",
    "ScreenBRank": "Screen B Rank",
    "Name": "Company",
    "PriceUSD": "Price (USD)",
    "MarketCapUSD": "Market Cap (USD)",
    "ADV20USD": "ADV20 (USD)",
    "Return3M": "Return 3M (%)",
    "RankReturn3M": "3M Sector PR",
    "EPSNow": "EPS FY1 Now",
    "EPS30d": "EPS FY1 30d Ago",
    "EPSRevision": "EPS Revision",
    "RankEPSRevision": "EPS Revision PR",
    "Price_t0": "Latest Local Price",
    "DistanceMA20": "Distance MA20",
    "RankDistanceMA20": "MA20 Sector PR",
    "DistanceMA50": "Distance MA50",
    "RankDistanceMA50": "MA50 Sector PR",
    "AvgVol5D": "Avg Volume 5D",
    "AvgVol20D": "Avg Volume 20D",
    "VolumeRatio5D20D": "Volume Ratio 5D/20D",
    "RankVolumeRatio": "Volume Sector PR",
    "PriceConfirmationA": "Price Confirmation A",
    "Return5D": "Return 5D (%)",
    "RankReturn5D": "5D Sector PR",
    "ScreenAScore": "Screen A Score",
    "NextEarningDate": "Next Earnings",
    "DaysToEarnings": "Business Days to Earnings",
    "CatalystScore": "Catalyst Score",
    "Vol20D": "Volatility 20D",
    "RankVol20D": "Volatility Sector PR",
    "ScreenBScore": "Screen B Score",
    "TurnoverRatio5D20D": "Turnover Ratio 5D/20D",
    "PriceHistoryCount": "Price Observations",
    "PassesOptionalFilter": "Optional Filter Flag",
}


def _select_and_rename(frame: pd.DataFrame, columns: list[str]) -> pd.DataFrame:
    output = frame.loc[:, columns].copy()
    return output.rename(columns={column: DISPLAY_NAMES.get(column, column) for column in columns})


def _methodology_table(result: ScreeningResult) -> pd.DataFrame:
    config = result.config
    rows = [
        (
            "Universe",
            "All valid tickers in the Main sheet",
            "No early economic filter is applied to either ranking.",
        ),
        (
            "Sector percentile",
            "(# lower + 0.5 x # equal) / N within sector",
            f"Missing observations receive {config.missing_percentile:.2f}.",
        ),
        (
            "3-month momentum",
            "Sector percentile of Bloomberg Return 3M %",
            f"Screen A weight: {config.a_return_3m_weight:.0%}.",
        ),
        (
            "5-day return",
            "Sector percentile of Bloomberg Return 5D %",
            f"Screen A weight: {config.a_return_5d_weight:.0%}; "
            f"Screen B weight: {config.b_return_5d_weight:.0%}.",
        ),
        (
            "EPS revision",
            "(EPS now - EPS 30d ago) / (|EPS now| + |EPS 30d ago| + epsilon)",
            f"Screen A weight: {config.a_eps_revision_weight:.0%}.",
        ),
        (
            "Volume ratio",
            "Average Volume 5D / Average Volume 20D",
            "Used in Screen A price confirmation and directly in Screen B.",
        ),
        (
            "Distance to MA20",
            "Latest local price / 20-day moving average - 1",
            "Requires 20 complete price observations.",
        ),
        (
            "Distance to MA50",
            "Latest local price / 50-day moving average - 1",
            "Requires 50 complete price observations.",
        ),
        (
            "Price confirmation A",
            f"{config.price_confirmation_ma20_weight:.0%} PR(D20) + "
            f"{config.price_confirmation_ma50_weight:.0%} PR(D50) + "
            f"{config.price_confirmation_volume_weight:.0%} PR(Volume ratio)",
            f"Screen A weight: {config.a_price_confirmation_weight:.0%}.",
        ),
        (
            "Screen A score",
            f"{config.a_return_3m_weight:.0%} PR(Return 3M) + "
            f"{config.a_eps_revision_weight:.0%} PR(EPS revision) + "
            f"{config.a_price_confirmation_weight:.0%} Price confirmation + "
            f"{config.a_return_5d_weight:.0%} PR(Return 5D)",
            "Core momentum ranking.",
        ),
        (
            "Days to earnings",
            f"Monday-Friday business days from {result.as_of_date:%Y-%m-%d}",
            "Local exchange holidays are not separately modeled.",
        ),
        (
            "Catalyst score",
            f"0-{config.catalyst_near_max_days} days: {config.catalyst_near_score:.2f}; "
            f"{config.catalyst_near_max_days + 1}-{config.catalyst_medium_max_days}: "
            f"{config.catalyst_medium_score:.2f}; "
            f"{config.catalyst_medium_max_days + 1}-{config.catalyst_far_max_days}: "
            f"{config.catalyst_far_score:.2f}; otherwise 0",
            f"Screen B weight: {config.b_catalyst_weight:.0%}.",
        ),
        (
            "20-day volatility",
            "Sector percentile of Bloomberg Vol 20D",
            f"Screen B weight: {config.b_volatility_weight:.0%}.",
        ),
        (
            "Screen B score",
            f"{config.b_catalyst_weight:.0%} Catalyst + "
            f"{config.b_volatility_weight:.0%} PR(Vol 20D) + "
            f"{config.b_return_5d_weight:.0%} PR(Return 5D) + "
            f"{config.b_volume_weight:.0%} PR(Volume ratio)",
            "Event-volatility ranking.",
        ),
        (
            "Optional diagnostic filter",
            f"Price >= ${config.optional_min_price_usd:g}; "
            f"market cap >= ${config.optional_min_market_cap_usd / 1e6:g}m; "
            f"ADV20 >= ${config.optional_min_adv20_usd / 1e6:g}m",
            "Displayed as a flag only; it never changes the rankings.",
        ),
    ]
    return pd.DataFrame(rows, columns=["Item", "Definition", "Use / Comment"])


def _quality_table(result: ScreeningResult) -> pd.DataFrame:
    return pd.DataFrame(
        [(key, value) for key, value in result.quality_summary.items()],
        columns=["Metric", "Value"],
    )


def _controls_table(result: ScreeningResult) -> pd.DataFrame:
    config = asdict(result.config)
    rows: list[tuple[str, Any]] = [
        ("Input file", result.input_path.name),
        ("Input path", str(result.input_path)),
        ("Model as-of date", result.as_of_date.strftime("%Y-%m-%d")),
    ]
    rows.extend((key, value) for key, value in config.items())
    return pd.DataFrame(rows, columns=["Control", "Value"])


def _format_workbook(path: Path, source_name: str) -> None:
    workbook = load_workbook(path)

    for sheet in workbook.worksheets:
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor=DARK_BLUE)
            cell.font = Font(color=WHITE, bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.row_dimensions[1].height = 32

        for column_cells in sheet.columns:
            header = str(column_cells[0].value or "")
            observed = [
                len(str(cell.value))
                for cell in column_cells[:200]
                if cell.value is not None
            ]
            width = max([len(header), *observed], default=10) + 2
            width = min(max(width, 10), 36)
            sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = width

        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=False)

    # Top sheets get score heat maps and source comments.
    for sheet_name in ["Screen A Top", "Screen B Top"]:
        sheet = workbook[sheet_name]
        score_header = "Screen A Score" if sheet_name == "Screen A Top" else "Screen B Score"
        score_column = next(
            cell.column
            for cell in sheet[1]
            if cell.value == score_header
        )
        letter = get_column_letter(score_column)
        sheet.conditional_formatting.add(
            f"{letter}2:{letter}{sheet.max_row}",
            ColorScaleRule(
                start_type="min",
                start_color="F8696B",
                mid_type="percentile",
                mid_value=50,
                mid_color="FFEB84",
                end_type="max",
                end_color="63BE7B",
            ),
        )
        for cell in sheet[1]:
            cell.comment = Comment(
                f"Source: {source_name}. Derived columns follow the Methodology sheet.",
                "OpenAI",
            )

    # Number formats for all sheets that contain rankings.
    for sheet_name in ["Screen A Top", "Screen B Top", "All Rankings"]:
        sheet = workbook[sheet_name]
        header_map = {cell.value: cell.column for cell in sheet[1]}

        def set_format(header: str, number_format: str) -> None:
            column = header_map.get(header)
            if column is None:
                return
            for row in range(2, sheet.max_row + 1):
                sheet.cell(row, column).number_format = number_format

        for header in [
            "3M Sector PR",
            "5D Sector PR",
            "EPS Revision PR",
            "MA20 Sector PR",
            "MA50 Sector PR",
            "Volume Sector PR",
            "Volatility Sector PR",
            "EPS Revision",
            "Distance MA20",
            "Distance MA50",
            "Price Confirmation A",
        ]:
            set_format(header, "0.0%")

        for header in ["Screen A Score", "Screen B Score", "Catalyst Score"]:
            set_format(header, "0.0000")

        for header in ["Volume Ratio 5D/20D", "Turnover Ratio 5D/20D"]:
            set_format(header, "0.00x")

        for header in ["Return 3M (%)", "Return 5D (%)", "Volatility 20D"]:
            set_format(header, "0.00")

        for header in ["Market Cap (USD)", "ADV20 (USD)"]:
            set_format(header, '$#,##0;[Red]($#,##0);-')

        set_format("Price (USD)", '$#,##0.00;[Red]($#,##0.00);-')
        set_format("Next Earnings", "dd-mmm-yyyy")

    # Make Methodology and Controls easier to read.
    for sheet_name in ["Methodology", "Data Quality", "Controls"]:
        sheet = workbook[sheet_name]
        sheet.freeze_panes = "A2"
        for row in range(2, sheet.max_row + 1):
            if row % 2 == 0:
                for cell in sheet[row]:
                    cell.fill = PatternFill("solid", fgColor=LIGHT_GRAY)
            for cell in sheet[row]:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    workbook.save(path)


def export_results(
    result: ScreeningResult,
    output_dir: str | Path,
    *,
    prefix: str = "WLS",
) -> dict[str, Path]:
    """Export one screening run to Excel, CSV, and JSON files."""

    destination = Path(output_dir).expanduser().resolve()
    destination.mkdir(parents=True, exist_ok=True)

    paths = {
        "excel": destination / f"{prefix}_Screen_A_B_Rankings.xlsx",
        "all_csv": destination / f"{prefix}_All_Rankings.csv",
        "screen_a_csv": destination / f"{prefix}_Screen_A_Top{result.config.top_n}.csv",
        "screen_b_csv": destination / f"{prefix}_Screen_B_Top{result.config.top_n}.csv",
        "metadata_json": destination / f"{prefix}_Run_Metadata.json",
    }

    # CSVs retain machine-readable column names.
    result.all_rankings.sort_values("ScreenARank").to_csv(paths["all_csv"], index=False)
    result.screen_a_top.sort_values("ScreenARank").to_csv(paths["screen_a_csv"], index=False)
    result.screen_b_top.sort_values("ScreenBRank").to_csv(paths["screen_b_csv"], index=False)

    metadata = {
        "as_of_date": result.as_of_date.strftime("%Y-%m-%d"),
        "input_path": str(result.input_path),
        "config": asdict(result.config),
        "quality_summary": result.quality_summary,
    }
    with paths["metadata_json"].open("w", encoding="utf-8") as handle:
        json.dump(metadata, handle, indent=2, ensure_ascii=False, default=str)

    top_a = _select_and_rename(
        result.screen_a_top.sort_values("ScreenARank"),
        TOP_A_COLUMNS,
    )
    top_b = _select_and_rename(
        result.screen_b_top.sort_values("ScreenBRank"),
        TOP_B_COLUMNS,
    )
    all_rankings = _select_and_rename(
        result.all_rankings.sort_values("ScreenARank"),
        ALL_COLUMNS,
    )

    with pd.ExcelWriter(paths["excel"], engine="openpyxl") as writer:
        top_a.to_excel(writer, sheet_name="Screen A Top", index=False)
        top_b.to_excel(writer, sheet_name="Screen B Top", index=False)
        all_rankings.to_excel(writer, sheet_name="All Rankings", index=False)
        _methodology_table(result).to_excel(writer, sheet_name="Methodology", index=False)
        _quality_table(result).to_excel(writer, sheet_name="Data Quality", index=False)
        _controls_table(result).to_excel(writer, sheet_name="Controls", index=False)

    _format_workbook(paths["excel"], result.input_path.name)
    return paths
