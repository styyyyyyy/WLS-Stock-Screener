"""Retrospective A3/B2 daily-close scenario for user-supplied WLS panels.

Inputs are local and never bundled with this repository. This is intentionally a
scenario, because an October membership, sector, earnings and liquidity snapshot
cannot reconstruct what was known on every September trading day.
"""

from __future__ import annotations

import argparse
import bisect
import csv
import json
import math
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path

import numpy as np
import pandas as pd
from dateutil.relativedelta import relativedelta
from openpyxl import load_workbook


PRICE_MIN_USD = 1.0
ADV20_MIN_USD = 15_000_000.0
MARKET_CAP_MIN_USD = 500_000_000.0


def numeric(value: object) -> float:
    try:
        result = float(value)
    except (TypeError, ValueError):
        return math.nan
    return result if math.isfinite(result) else math.nan


def parse_date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if value is None or not str(value).strip():
        return None
    parsed = pd.to_datetime(str(value), dayfirst=True, errors="coerce")
    return None if pd.isna(parsed) else parsed.date()


def load_panel(path: Path, sheet: str) -> tuple[list[date], list[str], np.ndarray, dict[str, dict]]:
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = book[sheet].iter_rows(values_only=True)
        header = next(rows)
        date_columns = sorted(
            (cell.date() if isinstance(cell, datetime) else cell, col)
            for col, cell in enumerate(header)
            if isinstance(cell, (date, datetime))
        )
        dates = [day for day, _ in date_columns]
        if not dates or len(set(dates)) != len(dates):
            raise ValueError(f"{sheet}: dates are absent or repeated")
        tickers: list[str] = []
        values: list[list[float]] = []
        metadata: dict[str, dict] = {}
        for row in rows:
            ticker = str(row[0] or "").strip()
            if not ticker:
                continue
            if ticker in metadata:
                raise ValueError(f"{sheet}: duplicate ticker {ticker}")
            tickers.append(ticker)
            values.append([numeric(row[col]) for _, col in date_columns])
            metadata[ticker] = {
                "name": str(row[1] or "").strip(),
                "currency": str(row[2] or "").strip() if sheet == "Prices" else "",
            }
        return dates, tickers, np.asarray(values, dtype=float), metadata
    finally:
        book.close()


def load_snapshot(path: Path, snapshot_date: date) -> tuple[dict[str, date | None], dict[str, float]]:
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = book["Worksheet"].iter_rows(values_only=True)
        header = next(rows)
        columns = {str(value).strip(): i for i, value in enumerate(header) if value is not None}
        needed = {"Ticker", "Expected Earnings", "Weight"}
        if not needed.issubset(columns):
            raise ValueError(f"Worksheet needs {sorted(needed)}")
        earnings: dict[str, date | None] = {}
        weights: dict[str, float] = {}
        for row in rows:
            ticker = str(row[columns["Ticker"]] or "").strip()
            if not ticker:
                continue
            if ticker in earnings:
                raise ValueError(f"Worksheet: duplicate ticker {ticker}")
            event = parse_date(row[columns["Expected Earnings"]])
            earnings[ticker] = event if event and event >= snapshot_date else None
            weights[ticker] = numeric(row[columns["Weight"]])
        return earnings, weights
    finally:
        book.close()


def load_sectors(path: Path) -> dict[str, str]:
    with path.open(newline="", encoding="utf-8-sig") as stream:
        reader = csv.DictReader(stream)
        if not reader.fieldnames or not {"Ticker", "Sector"}.issubset(reader.fieldnames):
            raise ValueError("Sector CSV needs Ticker,Sector")
        sectors: dict[str, str] = {}
        for row in reader:
            ticker = row["Ticker"].strip()
            if ticker in sectors:
                raise ValueError(f"Duplicate sector ticker {ticker}")
            sectors[ticker] = row["Sector"].strip()
        return sectors


def load_weight_snapshot(path: Path) -> dict[str, float]:
    """Optional later WLS weight snapshot already mapped to dated tickers."""
    with path.open(newline="", encoding="utf-8-sig") as stream:
        reader = csv.DictReader(stream)
        if not reader.fieldnames or not {"Ticker", "Weight"}.issubset(reader.fieldnames):
            raise ValueError("Weight CSV needs Ticker,Weight")
        weights: dict[str, float] = {}
        for row in reader:
            ticker = row["Ticker"].strip()
            if not ticker or ticker in weights:
                raise ValueError(f"Blank or duplicate weight ticker: {ticker!r}")
            weights[ticker] = numeric(row["Weight"])
        return weights


def load_filter_snapshot(path: Path, mode: str) -> tuple[set[str], dict]:
    """Read a *snapshot* cohort, never assume these values were known historically.

    ``eligible`` expects Ticker only, for an independently verified cohort.
    ``metrics`` expects Ticker,PriceD1USD,ADV20USD,MarketCapUSD and applies
    Bloomberg's strict > thresholds. Columns must be USD-converted upstream.
    """
    with path.open(newline="", encoding="utf-8-sig") as stream:
        reader = csv.DictReader(stream)
        fields = set(reader.fieldnames or [])
        needed = {"Ticker"} if mode == "eligible" else {
            "Ticker", "PriceD1USD", "ADV20USD", "MarketCapUSD"
        }
        if not needed.issubset(fields):
            raise ValueError(f"{mode} filter CSV needs columns {sorted(needed)}")
        seen: set[str] = set()
        eligible: set[str] = set()
        invalid = 0
        for row in reader:
            ticker = row["Ticker"].strip()
            if not ticker or ticker in seen:
                raise ValueError(f"Blank or duplicate filter ticker: {ticker!r}")
            seen.add(ticker)
            if mode == "eligible":
                eligible.add(ticker)
            else:
                price = numeric(row["PriceD1USD"])
                adv = numeric(row["ADV20USD"])
                cap = numeric(row["MarketCapUSD"])
                if not all(math.isfinite(x) for x in (price, adv, cap)):
                    invalid += 1
                if price > PRICE_MIN_USD and adv > ADV20_MIN_USD and cap > MARKET_CAP_MIN_USD:
                    eligible.add(ticker)
        return eligible, {
            "mode": f"{mode}_snapshot_proxy", "snapshot_rows": len(seen),
            "eligible_rows": len(eligible), "rows_with_missing_metrics": invalid,
            "strict_thresholds": {
                "price_d1_usd_gt": PRICE_MIN_USD,
                "adv20_usd_gt": ADV20_MIN_USD,
                "market_cap_usd_gt": MARKET_CAP_MIN_USD,
            },
        }


@dataclass
class Inputs:
    dates: list[date]
    tickers: list[str]
    prices: np.ndarray
    eps: np.ndarray
    volume: np.ndarray
    metadata: dict[str, dict]
    earnings: dict[str, date | None]
    weights: dict[str, float]
    sectors: dict[str, str]
    eligible: set[str]
    snapshot_date: date


def load_inputs(
    prices_path: Path, eps_volume_path: Path, sectors_path: Path,
    filter_path: Path, filter_mode: str, snapshot_date: date,
) -> tuple[Inputs, dict]:
    dates, tickers, prices, metadata = load_panel(prices_path, "Prices")
    eps_dates, eps_tickers, eps, _ = load_panel(eps_volume_path, "EPS")
    vol_dates, vol_tickers, volume, _ = load_panel(eps_volume_path, "Trad Volume")
    if dates != eps_dates or dates != vol_dates:
        raise ValueError("Price, EPS and volume date columns must match")
    if set(tickers) != set(eps_tickers) or set(tickers) != set(vol_tickers):
        raise ValueError("Price, EPS and volume ticker sets must match")
    eps_at = {ticker: i for i, ticker in enumerate(eps_tickers)}
    volume_at = {ticker: i for i, ticker in enumerate(vol_tickers)}
    eps = eps[[eps_at[ticker] for ticker in tickers], :]
    volume = volume[[volume_at[ticker] for ticker in tickers], :]
    earnings, weights = load_snapshot(prices_path, snapshot_date)
    if set(tickers) != set(earnings):
        raise ValueError("Worksheet and Prices ticker sets must match")
    sectors = load_sectors(sectors_path)
    missing_sector_rows = {
        ticker for ticker in tickers
        if sectors.get(ticker, "").strip().lower() in {"", "--", "unclassified"}
    }
    eligible, filter_audit = load_filter_snapshot(filter_path, filter_mode)
    unknown = eligible - set(tickers)
    if unknown:
        raise ValueError(f"Filter snapshot has {len(unknown)} tickers absent from Prices: {sorted(unknown)[:5]}")
    ranked_eligible = eligible - missing_sector_rows
    if len(ranked_eligible) < 5:
        raise ValueError("Fewer than five filtered stocks have a valid GICS sector")
    filter_audit.update({
        "panel_stocks": len(tickers), "eligible_present_in_panel": len(eligible),
        "eligible_with_valid_gics_sector": len(ranked_eligible),
        "eligible_without_gics_sector": len(eligible & missing_sector_rows),
        "warning": "This is a retrospective snapshot filter; historical daily price/ADV/cap and membership are unavailable.",
    })
    return Inputs(dates, tickers, prices, eps, volume, metadata, earnings, weights,
                  sectors, ranked_eligible, snapshot_date), filter_audit


def estimated_event(anchor: date | None, asof: date, snapshot_date: date) -> date | None:
    if anchor is None or anchor < snapshot_date:
        return None
    if anchor - relativedelta(months=3) >= snapshot_date:
        return anchor
    events = [anchor - relativedelta(months=3 * n) for n in range(0, 21)]
    return min((event for event in events if event >= asof), default=None)


def catalyst(event: date | None, asof: date) -> tuple[float, int | None]:
    if event is None:
        return 0.0, None
    days = int(np.busday_count(np.datetime64(asof), np.datetime64(event)))
    if 0 <= days <= 5:
        return 1.0, days
    if 6 <= days <= 10:
        return 0.75, days
    if 11 <= days <= 15:
        return 0.5, days
    return 0.0, days


def midpoint_percentile(series: pd.Series) -> pd.Series:
    valid = series.dropna()
    result = pd.Series(0.5, index=series.index, dtype=float)
    if len(valid):
        result.loc[valid.index] = (valid.rank(method="average") - 0.5) / len(valid)
    return result


class FeatureBank:
    def __init__(self, data: Inputs):
        self.data = data
        self.date_index = {day: j for j, day in enumerate(data.dates)}
        n, m = data.prices.shape
        self.active = (np.isfinite(data.prices) & (data.prices > 0)
                       & np.isfinite(data.volume) & (data.volume > 0))
        self.last_active = np.maximum.accumulate(
            np.where(self.active, np.arange(m)[None, :], -1), axis=1
        )
        self.return5 = np.full((n, m), np.nan)
        self.volume_ratio = np.full((n, m), np.nan)
        self.vol20 = np.full((n, m), np.nan)
        self.ma20 = np.full((n, m), np.nan)
        self.ma50 = np.full((n, m), np.nan)
        for row in range(n):
            active = np.flatnonzero(self.active[row])
            for k, j in enumerate(active):
                if k >= 5:
                    self.return5[row, j] = data.prices[row, j] / data.prices[row, active[k - 5]] - 1
                if k >= 19:
                    recent = active[k - 19:k + 1]
                    p = data.prices[row, recent]
                    v = data.volume[row, recent]
                    self.ma20[row, j] = p.mean()
                    self.volume_ratio[row, j] = v[-5:].mean() / v.mean() if v.mean() > 0 else np.nan
                    self.vol20[row, j] = np.std(np.diff(np.log(p)), ddof=1) * math.sqrt(260) * 100
                if k >= 49:
                    self.ma50[row, j] = data.prices[row, active[k - 49:k + 1]].mean()

    def ranked(self, signal: date) -> pd.DataFrame:
        """Compute features with signal-close observations and no later values."""
        data = self.data
        j = self.date_index[signal]
        # The original daily screener used the next global close as its model
        # as-of for calendar lookbacks, but never reads that future close's data.
        asof = data.dates[min(j + 1, len(data.dates) - 1)]
        rows = np.array([i for i, t in enumerate(data.tickers) if t in data.eligible], dtype=int)
        latest = self.last_active[rows, j]
        rows = rows[latest >= 0]
        latest = latest[latest >= 0]
        price = data.prices[rows, latest]
        old3_j = bisect.bisect_right(data.dates, asof - relativedelta(months=3)) - 1
        if old3_j >= 0:
            old_last = self.last_active[rows, old3_j]
            old_price = np.where(old_last >= 0, data.prices[rows, np.maximum(old_last, 0)], np.nan)
            with np.errstate(divide="ignore", invalid="ignore"):
                ret3 = price / old_price - 1
        else:
            ret3 = np.full(len(rows), np.nan)
        old_eps_j = bisect.bisect_right(data.dates, asof - timedelta(days=30)) - 1
        if old_eps_j >= 0:
            now, ago = data.eps[rows, j], data.eps[rows, old_eps_j]
            revision = (now - ago) / (np.abs(now) + np.abs(ago) + 1e-12)
        else:
            revision = np.full(len(rows), np.nan)
        with np.errstate(divide="ignore", invalid="ignore"):
            dist20 = price / self.ma20[rows, latest] - 1
            dist50 = price / self.ma50[rows, latest] - 1
        tickers = [data.tickers[i] for i in rows]
        frame = pd.DataFrame({
            "ticker": tickers,
            "sector": [data.sectors[t] for t in tickers],
            "active_on_signal": self.active[rows, j],
            "return3m": ret3, "return5d": self.return5[rows, latest],
            "eps_revision": revision, "distance20": dist20, "distance50": dist50,
            "volume_ratio": self.volume_ratio[rows, latest],
            "vol20": self.vol20[rows, latest],
        })
        for feature, rank in (
            ("return3m", "r3"), ("return5d", "r5"), ("eps_revision", "reps"),
            ("distance20", "rma20"), ("distance50", "rma50"),
            ("volume_ratio", "rvolume"), ("vol20", "rvol20"),
        ):
            frame[rank] = frame.groupby("sector")[feature].transform(midpoint_percentile)
        frame["confirmation"] = 0.40 * frame.rma20 + 0.30 * frame.rma50 + 0.30 * frame.rvolume
        frame["score_a"] = (0.40 * frame.r3 + 0.25 * frame.reps
                            + 0.25 * frame.confirmation + 0.10 * frame.r5)
        events = [estimated_event(data.earnings[t], asof, data.snapshot_date) for t in tickers]
        catalysts = [catalyst(event, asof) for event in events]
        frame["estimated_next_earnings"] = events
        frame["business_days_to_earnings"] = [item[1] for item in catalysts]
        frame["catalyst"] = [item[0] for item in catalysts]
        frame["score_b"] = (0.35 * frame.catalyst + 0.30 * frame.rvol20
                            + 0.20 * frame.r5 + 0.15 * frame.rvolume)
        return frame


def select_five(frame: pd.DataFrame) -> pd.DataFrame:
    a = frame.sort_values(
        ["score_a", "r3", "reps", "ticker"],
        ascending=[False, False, False, True], kind="mergesort"
    ).head(3).copy()
    b_ranked = frame.sort_values(
        ["score_b", "catalyst", "rvol20", "r5", "ticker"],
        ascending=[False, False, False, False, True], kind="mergesort"
    )
    b = b_ranked.loc[~b_ranked.ticker.isin(a.ticker)].head(2).copy()
    if len(a) != 3 or len(b) != 2:
        raise ValueError("Fewer than five distinct eligible stocks at this signal")
    a["screen"] = "A"
    b["screen"] = "B"
    selected = pd.concat([a, b], ignore_index=True)
    selected.insert(0, "slot", np.arange(1, 6))
    selected["target_weight"] = 0.20
    if selected.ticker.duplicated().any():
        raise AssertionError("A/B selection contains duplicate ticker")
    return selected


def execute(data: Inputs, bank: FeatureBank, selections: dict[date, list[str]],
            first_execution: date, final_mark: date) -> tuple[list[dict], list[dict]]:
    start = bank.date_index[first_execution]
    finish = bank.date_index[final_mark]
    row_of = {ticker: i for i, ticker in enumerate(data.tickers)}
    shares: dict[str, float] = {}
    cash = previous_nav = 1.0
    daily: list[dict] = []
    trades: list[dict] = []

    def mark(ticker: str, j: int) -> float:
        i = row_of[ticker]
        last = bank.last_active[i, j]
        if last < 0:
            raise ValueError(f"No price to mark {ticker} on {data.dates[j]}")
        return float(data.prices[i, last])

    for j in range(start, finish + 1):
        day = data.dates[j]
        signal = data.dates[j - 1]
        targets = selections.get(signal, []) if day != final_mark else []
        old = {t: q * mark(t, j) for t, q in shares.items()}
        nav = cash + sum(old.values())
        if not math.isfinite(nav) or nav <= 0:
            raise ValueError(f"Invalid NAV on {day}")
        realized_return = nav / previous_nav - 1
        open_targets = [t for t in targets if bank.active[row_of[t], j]]
        frozen = sum(value for t, value in old.items() if not bank.active[row_of[t], j])
        available = max(0.0, nav - frozen)
        desired_total = len(open_targets) * nav / 5
        scale = min(1.0, available / desired_total) if desired_total else 0.0
        desired = {t: nav / 5 * scale for t in open_targets}
        turnover = 0.0
        for ticker in sorted(set(old) | set(desired)):
            i = row_of[ticker]
            if not bank.active[i, j]:
                continue  # closed market: freeze shares, even if no longer selected
            before_value = old.get(ticker, 0.0)
            after_value = desired.get(ticker, 0.0)
            old_shares = shares.get(ticker, 0.0)
            new_shares = after_value / data.prices[i, j] if after_value > 1e-14 else 0.0
            change = new_shares - old_shares
            if abs(change) > 1e-14:
                trades.append({
                    "date": day.isoformat(), "signal_date": signal.isoformat() if day != final_mark else "",
                    "ticker": ticker, "side": "BUY" if change > 0 else "SELL",
                    "shares_change": change, "execution_close": float(data.prices[i, j]),
                    "value_change": after_value - before_value,
                    "reason": "terminal_liquidation" if day == final_mark else "daily_rebalance",
                })
            turnover += abs(after_value - before_value) / nav
            if new_shares > 1e-14:
                shares[ticker] = new_shares
            else:
                shares.pop(ticker, None)
        cash = nav - sum(q * mark(t, j) for t, q in shares.items())
        if cash < -1e-9:
            raise ValueError(f"Negative cash on {day}: {cash}")
        cash = max(cash, 0.0)
        daily.append({
            "date": day.isoformat(), "signal_date": signal.isoformat() if day != final_mark else "",
            "nav": nav, "daily_return": realized_return, "cash_after_rebalance": cash,
            "positions_after_rebalance": len(shares),
            "frozen_positions_after_rebalance": sum(not bank.active[row_of[t], j] for t in shares),
            "target_slots": len(targets),
            "closed_target_slots": len(targets) - len(open_targets),
            "gross_turnover": turnover,
        })
        previous_nav = nav
    return daily, trades


def weighted_wls_proxy(data: Inputs, bank: FeatureBank, first_execution: date,
                       final_mark: date) -> list[dict]:
    """Retrospective buy-and-hold proxy from *full* WLS snapshot weights."""
    start = bank.date_index[first_execution]
    finish = bank.date_index[final_mark]
    price_valid = np.isfinite(data.prices) & (data.prices > 0)
    valid = []
    for i, ticker in enumerate(data.tickers):
        weight = data.weights.get(ticker, math.nan)
        if not math.isfinite(weight) or weight <= 0:
            continue
        valid.append((i, weight))
    if not valid:
        raise ValueError("No positive WLS snapshot weights")
    total_weight = sum(weight for _, weight in valid)
    # A member without an initial close is held flat, rather than silently
    # reallocating its weight among priced constituents.
    units = [(i, weight / total_weight, float(data.prices[i, start])
              if price_valid[i, start] else math.nan)
             for i, weight in valid]
    latest_price = np.maximum.accumulate(
        np.where(price_valid, np.arange(len(data.dates))[None, :], -1), axis=1
    )
    return [{
        "date": data.dates[j].isoformat(),
        "snapshot_weighted_proxy_nav": sum(
            weight * (data.prices[i, latest_price[i, j]] / start_price
                      if math.isfinite(start_price) and latest_price[i, j] >= 0 else 1.0)
            for i, weight, start_price in units
        ),
        "constituent_count": len(units),
        "unpriced_at_start_count": sum(not math.isfinite(start_price) for _, _, start_price in units),
    } for j in range(start, finish + 1)]


def write_csv(path: Path, records: list[dict]) -> None:
    if not records:
        raise ValueError(f"No rows to write: {path}")
    columns = list(dict.fromkeys(key for row in records for key in row))
    with path.open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.DictWriter(stream, fieldnames=columns)
        writer.writeheader()
        writer.writerows(records)


def run(data: Inputs, filter_audit: dict, first_signal: date,
        final_mark: date) -> tuple[dict, list[dict], list[dict], list[dict], list[dict]]:
    bank = FeatureBank(data)
    if first_signal not in bank.date_index or final_mark not in bank.date_index:
        raise ValueError("Signal or final mark date absent from panel")
    first_j, final_j = bank.date_index[first_signal], bank.date_index[final_mark]
    if final_j - first_j < 2:
        raise ValueError("Need at least one execution and a later terminal mark")
    selections: dict[date, list[str]] = {}
    holdings: list[dict] = []
    for j in range(first_j, final_j - 1):
        signal, execution_day = data.dates[j], data.dates[j + 1]
        top = select_five(bank.ranked(signal))
        selections[signal] = top.ticker.tolist()
        for item in top.to_dict("records"):
            ticker = item["ticker"]
            holdings.append({
                "signal_date": signal.isoformat(), "execution_date": execution_day.isoformat(),
                "slot": int(item["slot"]), "screen": item["screen"], "ticker": ticker,
                "name": data.metadata[ticker]["name"], "sector": item["sector"],
                "score": float(item["score_a"] if item["screen"] == "A" else item["score_b"]),
                "target_weight": 0.20, "return3m": float(item["return3m"]),
                "return5d": float(item["return5d"]),
                "eps_revision": float(item["eps_revision"]),
                "vol20_pct": float(item["vol20"]),
                "volume_ratio": float(item["volume_ratio"]),
                "earnings_anchor_snapshot": data.earnings[ticker].isoformat()
                if data.earnings[ticker] else "",
                "estimated_next_earnings": item["estimated_next_earnings"].isoformat()
                if isinstance(item["estimated_next_earnings"], date) else "",
                "business_days_to_earnings": item["business_days_to_earnings"],
                "catalyst": float(item["catalyst"]),
                "active_on_signal": bool(item["active_on_signal"]),
                "active_on_execution": bool(bank.active[data.tickers.index(ticker), j + 1]),
            })
    first_execution = data.dates[first_j + 1]
    daily, trades = execute(data, bank, selections, first_execution, final_mark)
    reference = weighted_wls_proxy(data, bank, first_execution, final_mark)
    navs = np.array([1.0] + [row["nav"] for row in daily])
    peaks = np.maximum.accumulate(navs)
    stats = {
        "gross_strategy_return_pct": (navs[-1] - 1) * 100,
        "max_close_to_close_drawdown_pct": min(navs / peaks - 1) * 100,
        "snapshot_weighted_wls_proxy_return_pct": (reference[-1]["snapshot_weighted_proxy_nav"] - 1) * 100,
        "return_minus_snapshot_wls_proxy_percentage_points":
            (navs[-1] - reference[-1]["snapshot_weighted_proxy_nav"]) * 100,
        "signal_days": len(selections), "selected_stock_days": len(holdings),
        "executed_buys_sells": len(trades),
        "selected_unfilled_at_next_close": sum(not row["active_on_execution"] for row in holdings),
        "gross_turnover_sum": sum(row["gross_turnover"] for row in daily),
        "final_unliquidated_positions": daily[-1]["positions_after_rebalance"],
    }
    summary = {
        "period": {"first_signal": first_signal.isoformat(),
                   "first_execution": first_execution.isoformat(),
                   "last_execution": data.dates[final_j - 1].isoformat(),
                   "terminal_mark": final_mark.isoformat()},
        "filter": filter_audit,
        "method": "A top 3 + distinct B top 2; equal 20% NAV targets; daily next-global-close rebalance; carry overnight; closed-market shares frozen and unavailable slots cash; midpoint percentiles within GICS sector.",
        "benchmark": "Full-universe WLS snapshot-weighted local-price buy-and-hold proxy; not official historical WLS Index or alpha.",
        "limitations": [
            "Current membership, GICS sectors, liquidity/cap filter and earnings snapshot are applied retrospectively; this is not a point-in-time backtest.",
            "Historical reproduction convention: three-month and EPS 30-calendar-day reference dates, plus earnings countdown, are measured from the following global close; all price, EPS and volume observations used for the signal are dated on or before the signal close.",
            "Earnings dates are inferred by moving snapshot next-earnings dates backward in three-calendar-month steps.",
            "Prices are compounded as local-currency ratios without FX, dividends, verified corporate-action adjustments or trading costs.",
            "No official historical WLS index levels are used; the WLS comparator uses future snapshot weights.",
        ],
        "statistics": stats,
    }
    return summary, daily, holdings, trades, reference


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--prices", required=True, type=Path, help="WLS workbook with Worksheet and Prices")
    parser.add_argument("--eps-volume", required=True, type=Path, help="Workbook with EPS and Trad Volume")
    parser.add_argument("--sectors", required=True, type=Path, help="Ticker,Sector CSV mapped to dated tickers")
    filter_group = parser.add_mutually_exclusive_group(required=True)
    filter_group.add_argument("--eligible-snapshot", type=Path, help="Verified snapshot cohort CSV: Ticker")
    filter_group.add_argument("--filter-snapshot", type=Path,
                              help="Snapshot CSV: Ticker,PriceD1USD,ADV20USD,MarketCapUSD")
    parser.add_argument("--snapshot-date", required=True, type=date.fromisoformat)
    parser.add_argument("--wls-weights-snapshot", type=Path,
                        help="Optional full-WLS Ticker,Weight CSV mapped to dated tickers")
    parser.add_argument("--first-signal", type=date.fromisoformat, default=date(2026, 9, 7))
    parser.add_argument("--final-mark", type=date.fromisoformat, default=date(2026, 10, 7))
    parser.add_argument("--output-dir", required=True, type=Path)
    args = parser.parse_args()
    filter_path = args.eligible_snapshot or args.filter_snapshot
    mode = "eligible" if args.eligible_snapshot else "metrics"
    inputs, filter_audit = load_inputs(
        args.prices, args.eps_volume, args.sectors, filter_path, mode, args.snapshot_date
    )
    if args.wls_weights_snapshot:
        weights = load_weight_snapshot(args.wls_weights_snapshot)
        if set(weights) != set(inputs.tickers):
            raise ValueError("WLS weight snapshot must cover the full dated panel exactly")
        inputs.weights = weights
    summary, daily, holdings, trades, reference = run(
        inputs, filter_audit, args.first_signal, args.final_mark
    )
    args.output_dir.mkdir(parents=True, exist_ok=True)
    (args.output_dir / "summary.json").write_text(
        json.dumps(summary, indent=2, allow_nan=False) + "\n", encoding="utf-8"
    )
    write_csv(args.output_dir / "daily.csv", daily)
    write_csv(args.output_dir / "holdings.csv", holdings)
    write_csv(args.output_dir / "trades.csv", trades)
    write_csv(args.output_dir / "wls_snapshot_proxy_daily.csv", reference)
    print(json.dumps(summary["statistics"], indent=2))


if __name__ == "__main__":
    main()
